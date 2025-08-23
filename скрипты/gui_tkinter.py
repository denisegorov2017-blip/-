import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import threading
import os
import pandas as pd
import webbrowser
import json
try:
    import PIL.Image as Image
    import PIL.ImageDraw as ImageDraw
    HAS_PIL = True
except ImportError:
    HAS_PIL = False
from improved_coefficient_calculator import main as calc_main
from analytics import forecast_shrinkage, compare_coefficients, cluster_nomenclatures

class ShrinkageCalculatorGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Расчет коэффициентов усушки")
        self.root.geometry("1100x750")
        self.root.minsize(900, 600)
        
        # Создаем иконку приложения
        self.create_app_icon()
        
        # Настройка стилей
        self.setup_styles()
        
        # Получаем пути к файлам
        self.script_dir = os.path.dirname(os.path.abspath(__file__))
        self.project_root = os.path.dirname(self.script_dir)
        self.default_input_file = os.path.join(self.project_root, "исходные_данные", "sheet_1_Лист_1.csv")
        self.csv_output_file = os.path.join(self.project_root, "результаты", "коэффициенты_усушки_улучшенные.csv")
        self.html_output_file = os.path.join(self.project_root, "результаты", "коэффициенты_усушки_улучшенные.html")
        self.config_file = os.path.join(self.project_root, "config.json")
        self.icon_file = os.path.join(self.project_root, "icon.ico")
        
        # Загружаем настройки
        self.load_config()
        
        self.results_data = None
        self.create_widgets()
        self.load_last_session()
        
    def create_app_icon(self):
        """Создание иконки приложения"""
        if HAS_PIL:
            try:
                # Создаем изображение для иконки
                size = 32
                image = Image.new('RGBA', (size, size), (74, 144, 226, 255))  # Синий цвет
                draw = ImageDraw.Draw(image)
                
                # Рисуем упрощенное изображение графика
                draw.line([(5, 25), (15, 15), (25, 20)], fill=(255, 255, 255), width=2)
                draw.ellipse([(22, 17), (28, 23)], fill=(255, 255, 255))
                
                # Сохраняем иконку
                if not os.path.exists(self.icon_file):
                    image.save(self.icon_file, format='ICO')
                
                # Устанавливаем иконку для окна
                self.root.iconbitmap(self.icon_file)
            except Exception:
                # Если не удалось создать иконку, используем стандартную
                pass
        else:
            # Если PIL не установлен, используем стандартную иконку
            try:
                self.root.iconbitmap(default='@icon')
            except Exception:
                pass
                
    def setup_styles(self):
        """Настройка стилей для виджетов"""
        style = ttk.Style()
        
        # Цветовая схема с темным текстом для лучшей видимости
        bg_color = "#f0f0f0"
        accent_color = "#BBDEFB"  # Светло-синий фон
        success_color = "#C8E6C9"  # Светло-зеленый фон
        warning_color = "#FFE0B2"  # Светло-оранжевый фон
        error_color = "#FFCDD2"    # Светло-красный фон
        
        # Стиль для прогресс-бара
        style.configure("Custom.Horizontal.TProgressbar", 
                       thickness=25,  # Более толстый прогресс-бар
                       background="#2196F3",  # Ярко-синий цвет
                       troughcolor="#E0E0E0",  # Цвет фона
                       bordercolor="#BDBDBD",  # Цвет границы
                       lightcolor="#64B5F6",  # Светлый цвет
                       darkcolor="#1976D2")   # Темный цвет
        
        # Стиль для кнопок с темным текстом
        style.configure("Accent.TButton", 
                       background=accent_color, 
                       foreground="#0D47A1",  # Темно-синий текст
                       font=("Arial", 9, "bold"),
                       padding=6)
        
        style.map("Accent.TButton", 
                 background=[("active", "#90CAF9")])  # Более темный при наведении
        
        style.configure("Success.TButton", 
                       background=success_color, 
                       foreground="#1B5E20",  # Темно-зеленый текст
                       font=("Arial", 9, "bold"),
                       padding=6)
        
        style.map("Success.TButton", 
                 background=[("active", "#A5D6A7")])
        
        style.configure("Warning.TButton", 
                       background=warning_color, 
                       foreground="#E65100",  # Темно-оранжевый текст
                       font=("Arial", 9, "bold"),
                       padding=6)
        
        style.map("Warning.TButton", 
                 background=[("active", "#FFCC80")])
        
        style.configure("Error.TButton", 
                       background=error_color, 
                       foreground="#B71C1C",  # Темно-красный текст
                       font=("Arial", 9, "bold"),
                       padding=6)
        
        style.map("Error.TButton", 
                 background=[("active", "#EF9A9A")])
        
        # Стиль для заголовков
        style.configure("Title.TLabel", 
                       font=("Arial", 12, "bold"),
                       foreground="#212121")
        
        # Стиль для фреймов
        style.configure("Card.TFrame", 
                       background="white")
        
        style.configure("TLabelframe", 
                       background=bg_color,
                       foreground="#212121")
        
        style.configure("TLabelframe.Label", 
                       background=bg_color,
                       foreground="#212121",
                       font=("Arial", 10, "bold"))
        
    def create_widgets(self):
        # Главный фрейм
        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.grid(row=0, column=0, sticky=(tk.N, tk.S, tk.E, tk.W))
        
        # Заголовок
        title_label = ttk.Label(main_frame, text="Система расчета коэффициентов усушки", style="Title.TLabel")
        title_label.grid(row=0, column=0, columnspan=3, sticky=tk.W, pady=(0, 15))
        
        # Верхняя панель
        top_frame = ttk.LabelFrame(main_frame, text="Настройки", padding="10")
        top_frame.grid(row=1, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 15))
        
        # Выбор файла
        ttk.Label(top_frame, text="Файл с данными:").grid(row=0, column=0, sticky=tk.W, pady=(0, 5))
        file_frame = ttk.Frame(top_frame)
        file_frame.grid(row=1, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=(0, 10))
        
        self.file_path_var = tk.StringVar(value=self.default_input_file)
        self.file_entry = ttk.Entry(file_frame, textvariable=self.file_path_var, width=70)
        self.file_entry.grid(row=0, column=0, sticky=(tk.W, tk.E), padx=(0, 5))
        file_frame.columnconfigure(0, weight=1)
        
        ttk.Button(file_frame, text="Обзор", command=self.browse_file, style="Accent.TButton").grid(row=0, column=1)
        
        # Кнопки управления
        button_frame = ttk.Frame(top_frame)
        button_frame.grid(row=2, column=0, columnspan=3, pady=(10, 0), sticky=(tk.W, tk.E))
        
        self.calc_button = ttk.Button(button_frame, text="Рассчитать коэффициенты", 
                                     command=self.start_calculation, style="Accent.TButton")
        self.calc_button.pack(side=tk.LEFT, padx=(0, 5))
        
        self.view_button = ttk.Button(button_frame, text="Просмотр результатов", 
                                     command=self.view_results, state=tk.DISABLED, style="Accent.TButton")
        self.view_button.pack(side=tk.LEFT, padx=(0, 5))
        
        self.html_button = ttk.Button(button_frame, text="Открыть HTML", 
                                     command=self.open_html, state=tk.DISABLED, style="Accent.TButton")
        self.html_button.pack(side=tk.LEFT, padx=(0, 5))
        
        self.unprocessed_button = ttk.Button(button_frame, text="Необработанные позиции", 
                                           command=self.view_unprocessed, state=tk.DISABLED, style="Warning.TButton")
        self.unprocessed_button.pack(side=tk.LEFT, padx=(0, 5))
        
        self.forecast_button = ttk.Button(button_frame, text="Прогноз усушки", command=self.forecast_shrinkage, state=tk.DISABLED, style="Accent.TButton")
        self.forecast_button.pack(side=tk.LEFT, padx=(0, 5))
        
        self.compare_button = ttk.Button(button_frame, text="Сравнить коэффициенты", command=self.compare_coefficients, style="Accent.TButton")
        self.compare_button.pack(side=tk.LEFT, padx=(0, 5))
        
        self.cluster_button = ttk.Button(button_frame, text="Кластеризация", command=self.cluster_nomenclatures, state=tk.DISABLED, style="Accent.TButton")
        self.cluster_button.pack(side=tk.LEFT, padx=(0, 5))
        
        ttk.Button(button_frame, text="Экспорт", command=self.export_results, style="Warning.TButton").pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(button_frame, text="Выход", command=self.on_exit, style="Error.TButton").pack(side=tk.RIGHT)
        
        # Прогресс бар и статус
        progress_frame = ttk.Frame(top_frame)
        progress_frame.grid(row=3, column=0, columnspan=3, pady=(10, 0), sticky=(tk.W, tk.E))
        
        # Определение стиля для прогресс-бара
        style = ttk.Style()
        style.configure("Custom.Horizontal.TProgressbar", 
                       thickness=20,  # Более толстый прогресс-бар
                       background="#4a90e2")  # Синий цвет для лучшей видимости
        
        self.progress = ttk.Progressbar(progress_frame, mode='indeterminate', style="Custom.Horizontal.TProgressbar")
        self.progress.grid(row=0, column=0, sticky=(tk.W, tk.E), padx=(0, 10))
        progress_frame.columnconfigure(0, weight=1)
        
        # Текстовое отображение процента выполнения
        self.progress_percent_var = tk.StringVar(value="")
        self.progress_percent_label = ttk.Label(progress_frame, textvariable=self.progress_percent_var, foreground="blue", font=("Arial", 9, "bold"))
        self.progress_percent_label.grid(row=0, column=1, padx=(0, 10))
        
        self.status_var = tk.StringVar(value="Готов к работе")
        self.status_label = ttk.Label(progress_frame, textvariable=self.status_var, foreground="gray")
        self.status_label.grid(row=0, column=2)
        
        # Вывод логов
        log_frame = ttk.LabelFrame(main_frame, text="Логи выполнения", padding="10")
        log_frame.grid(row=2, column=0, columnspan=3, sticky=(tk.N, tk.S, tk.E, tk.W), pady=(0, 15))
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)
        
        self.log_text = scrolledtext.ScrolledText(log_frame, height=8, wrap=tk.WORD)
        self.log_text.grid(row=0, column=0, sticky=(tk.N, tk.S, tk.E, tk.W))
        
        # Поиск и таблица результатов
        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.grid(row=3, column=0, columnspan=3, sticky=(tk.N, tk.S, tk.E, tk.W))
        bottom_frame.columnconfigure(0, weight=1)
        bottom_frame.rowconfigure(1, weight=1)
        
        # Поиск
        search_frame = ttk.Frame(bottom_frame)
        search_frame.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 10))
        search_frame.columnconfigure(1, weight=1)
        
        ttk.Label(search_frame, text="Поиск:").pack(side=tk.LEFT)
        self.search_var = tk.StringVar()
        self.search_entry = ttk.Entry(search_frame, textvariable=self.search_var, width=30)
        self.search_entry.pack(side=tk.LEFT, padx=(5, 5))
        ttk.Button(search_frame, text="Найти", command=self.search_results).pack(side=tk.LEFT)
        
        # Таблица результатов
        table_frame = ttk.LabelFrame(bottom_frame, text="Результаты расчета", padding="5")
        table_frame.grid(row=1, column=0, sticky=(tk.N, tk.S, tk.E, tk.W), pady=(0, 10))
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        
        # Создаем Treeview для отображения результатов
        columns = ("Номенклатура", "Коэф. A", "Коэф. B", "Точность", "Усушка (кг)", "Дата")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=12)
        
        # Определяем заголовки и ширину колонок
        column_widths = [250, 80, 80, 80, 100, 100]
        for i, col in enumerate(columns):
            self.tree.heading(col, text=col)
            self.tree.column(col, width=column_widths[i], anchor=tk.CENTER)
        
        # Создаем скроллбары
        v_scrollbar = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        h_scrollbar = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scrollbar.set, xscrollcommand=h_scrollbar.set)
        
        # Размещаем элементы
        self.tree.grid(row=0, column=0, sticky=(tk.N, tk.S, tk.E, tk.W))
        v_scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
        h_scrollbar.grid(row=1, column=0, sticky=(tk.W, tk.E))
        
        # Настраиваем весовые коэффициенты для растягивания
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(2, weight=1)
        main_frame.rowconfigure(3, weight=2)
        
        # Привязываем событие изменения текста в поле поиска
        self.search_var.trace("w", self.on_search_change)
        
        # Привязываем событие закрытия окна
        self.root.protocol("WM_DELETE_WINDOW", self.on_exit)
        
    def load_config(self):
        """Загрузка конфигурации из файла"""
        self.config = {
            "last_input_file": self.default_input_file,
            "window_geometry": "1100x750"
        }
        
        if os.path.exists(self.config_file):
            try:
                with open(self.config_file, 'r', encoding='utf-8') as f:
                    self.config.update(json.load(f))
            except Exception as e:
                self.log_message(f"Ошибка загрузки конфигурации: {str(e)}")
                
    def save_config(self):
        """Сохранение конфигурации в файл"""
        self.config["last_input_file"] = self.file_path_var.get()
        self.config["window_geometry"] = self.root.geometry()
        
        try:
            with open(self.config_file, 'w', encoding='utf-8') as f:
                json.dump(self.config, f, ensure_ascii=False, indent=2)
        except Exception as e:
            self.log_message(f"Ошибка сохранения конфигурации: {str(e)}")
            
    def load_last_session(self):
        """Загрузка последней сессии"""
        if "last_input_file" in self.config:
            self.file_path_var.set(self.config["last_input_file"])
            
        if "window_geometry" in self.config:
            self.root.geometry(self.config["window_geometry"])
            
    def browse_file(self):
        filename = filedialog.askopenfilename(
            title="Выберите файл с данными",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")]
        )
        if filename:
            self.file_path_var.set(filename)
            
    def log_message(self, message):
        """Добавление сообщения в лог"""
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.root.update_idletasks()
        
    def update_status(self, message, color="gray"):
        """Обновление статуса"""
        self.status_var.set(message)
        self.status_label.configure(foreground=color)
        
    def start_calculation(self):
        """Запуск расчета в отдельном потоке"""
        # Проверяем существование файла
        if not os.path.exists(self.file_path_var.get()):
            messagebox.showerror("Ошибка", "Указанный файл не существует")
            return
            
        # Запускаем расчет в отдельном потоке, чтобы не блокировать GUI
        thread = threading.Thread(target=self.calculate)
        thread.daemon = True
        thread.start()
        
    def calculate(self):
        """Выполнение расчета"""
        try:
            self.calc_button.configure(state=tk.DISABLED)
            # Устанавливаем режим определенного прогресса
            self.progress.configure(mode='determinate', maximum=100)
            self.progress['value'] = 0
            self.progress_percent_var.set("0%")
            self.update_status("Выполняется расчет...", "blue")
            self.log_message("=" * 50)
            self.log_message("Начинаем расчет коэффициентов усушки...")
            
            # Запускаем расчет
            calc_main()
            
            # Имитация прогресса (если основной расчет не поддерживает callback)
            for i in range(101):
                self.progress['value'] = i
                self.progress_percent_var.set(f"{i}%")
                self.root.update_idletasks()
                # Небольшая задержка для демонстрации прогресса
                self.root.after(20)
            
            # Загружаем результаты
            if os.path.exists(self.csv_output_file):
                self.results_data = pd.read_csv(self.csv_output_file)
                self.update_results_table(self.results_data)
                self.log_message("Расчет успешно завершен!")
                self.update_status("Расчет завершен успешно", "green")
                self.progress_percent_var.set("100%")
                
                # Активируем кнопки результатов
                self.view_button.configure(state=tk.NORMAL)
                self.html_button.configure(state=tk.NORMAL)
                self.unprocessed_button.configure(state=tk.NORMAL)
                self.forecast_button.configure(state=tk.NORMAL)
                self.cluster_button.configure(state=tk.NORMAL)
            else:
                self.log_message("Ошибка: файл с результатами не найден")
                self.update_status("Ошибка при расчете", "red")
                self.progress_percent_var.set("Ошибка")
                
        except Exception as e:
            self.log_message(f"Ошибка: {str(e)}")
            self.update_status("Ошибка при расчете", "red")
            self.progress_percent_var.set("Ошибка")
        finally:
            self.calc_button.configure(state=tk.NORMAL)
            
    def update_results_table(self, data):
        """Обновление таблицы результатов"""
        # Очищаем таблицу
        for item in self.tree.get_children():
            self.tree.delete(item)
            
        # Заполняем таблицу данными
        for _, row in data.iterrows():
            # Извлекаем значение усушки из примечания
            shrinkage = "н/д"
            if 'Усушка' in str(row['Примечание']):
                try:
                    shrinkage = f"{float(row['Примечание'].split('Усушка ')[1].split(' кг')[0]):.3f}"
                except:
                    pass
                    
            self.tree.insert("", "end", values=(
                row["Номенклатура"],
                f"{row['a']:.3f}",
                f"{row['b (день⁻¹)']:.3f}",
                f"{row['Точность (%)']:.1f}%",
                shrinkage,
                row['Дата_расчета']
            ))
            
    def view_results(self):
        """Просмотр полных результатов"""
        if self.results_data is not None:
            # Открываем окно с полными результатами
            results_window = tk.Toplevel(self.root)
            results_window.title("Полные результаты расчета")
            results_window.geometry("900x700")
            results_window.minsize(700, 500)
            
            # Создаем текстовое поле для отображения результатов
            text_frame = ttk.Frame(results_window, padding="10")
            text_frame.pack(expand=True, fill='both')
            
            text_area = scrolledtext.ScrolledText(text_frame, wrap=tk.WORD)
            text_area.pack(expand=True, fill='both')
            
            # Формируем текст с результатами
            result_text = "ПОЛНЫЕ РЕЗУЛЬТАТЫ РАСЧЕТА КОЭФФИЦИЕНТОВ УСУШКИ\n"
            result_text += "=" * 60 + "\n\n"
            
            for _, row in self.results_data.iterrows():
                result_text += f"Номенклатура: {row['Номенклатура']}\n"
                result_text += f"Коэффициент A: {row['a']:.3f}\n"
                result_text += f"Коэффициент B: {row['b (день⁻¹)']:.3f}\n"
                result_text += f"Коэффициент C: {row['c']:.3f}\n"
                result_text += f"Точность: {row['Точность (%)']:.1f}%\n"
                result_text += f"Дата расчета: {row['Дата_расчета']}\n"
                result_text += f"Примечание: {row['Примечание']}\n"
                result_text += "-" * 50 + "\n\n"
                
            text_area.insert(tk.END, result_text)
            text_area.config(state=tk.DISABLED)
        else:
            messagebox.showwarning("Предупреждение", "Результаты еще не рассчитаны")
            
    def open_html(self):
        """Открытие HTML-файла с результатами"""
        if os.path.exists(self.html_output_file):
            try:
                webbrowser.open(f'file://{os.path.abspath(self.html_output_file)}')
                self.log_message("Открываем HTML-файл в браузере...")
            except Exception as e:
                self.log_message(f"Ошибка открытия HTML: {str(e)}")
                messagebox.showerror("Ошибка", f"Не удалось открыть HTML-файл:\n{str(e)}")
        else:
            messagebox.showwarning("Предупреждение", "HTML-файл с результатами не найден")
            
    def view_unprocessed(self):
        """Открытие HTML-файла с необработанными позициями"""
        unprocessed_html_file = os.path.join(self.project_root, "результаты", "необработанные_позиции.html")
        if os.path.exists(unprocessed_html_file):
            try:
                webbrowser.open(f'file://{os.path.abspath(unprocessed_html_file)}')
                self.log_message("Открываем файл с необработанными позициями в браузере...")
            except Exception as e:
                self.log_message(f"Ошибка открытия файла с необработанными позициями: {str(e)}")
                messagebox.showerror("Ошибка", f"Не удалось открыть файл с необработанными позициями:\n{str(e)}")
        else:
            messagebox.showwarning("Предупреждение", "Файл с необработанными позициями не найден")
            
    def export_results(self):
        """Экспорт результатов"""
        if self.results_data is None:
            messagebox.showwarning("Предупреждение", "Нет результатов для экспорта")
            return
            
        # Проверяем наличие библиотеки openpyxl для экспорта в Excel
        try:
            import openpyxl
            has_openpyxl = True
        except ImportError:
            has_openpyxl = False
            
        # Определяем доступные форматы
        filetypes = [("CSV files", "*.csv")]
        if has_openpyxl:
            filetypes.insert(0, ("Excel files", "*.xlsx"))
        filetypes.append(("Text files", "*.txt"))
        filetypes.append(("All files", "*.*"))
            
        # Выбираем место для сохранения
        filename = filedialog.asksaveasfilename(
            title="Сохранить результаты",
            defaultextension=".csv",
            filetypes=filetypes
        )
        
        if filename:
            try:
                if filename.endswith('.xlsx') and has_openpyxl:
                    self.results_data.to_excel(filename, index=False, engine='openpyxl')
                elif filename.endswith('.txt'):
                    with open(filename, 'w', encoding='utf-8') as f:
                        self.results_data.to_string(f, index=False)
                else:
                    self.results_data.to_csv(filename, index=False, encoding='utf-8')
                    
                self.log_message(f"Результаты экспортированы в: {filename}")
                messagebox.showinfo("Успех", f"Результаты успешно экспортированы в:\n{filename}")
            except Exception as e:
                self.log_message(f"Ошибка экспорта: {str(e)}")
                messagebox.showerror("Ошибка", f"Не удалось экспортировать результаты:\n{str(e)}")
                
    def on_search_change(self, *args):
        """Обработчик изменения текста поиска"""
        self.search_results()
        
    def search_results(self):
        """Поиск по результатам"""
        if self.results_data is None:
            return
            
        search_term = self.search_var.get().lower()
        if not search_term:
            # Если строка поиска пуста, показываем все результаты
            filtered_data = self.results_data
        else:
            # Фильтруем данные по названию номенклатуры
            filtered_data = self.results_data[
                self.results_data["Номенклатура"].str.lower().str.contains(search_term, na=False)
            ]
            
        self.update_results_table(filtered_data)
        
    def on_exit(self):
        """Обработчик выхода из приложения"""
        # Сохраняем конфигурацию
        self.save_config()
        
        # Закрываем приложение
        self.root.quit()
        
    def forecast_shrinkage(self):
        """Прогнозирование усушки"""
        if self.results_data is None:
            messagebox.showwarning("Предупреждение", "Результаты еще не рассчитаны")
            return
            
        # Открываем окно для ввода параметров прогноза
        forecast_window = tk.Toplevel(self.root)
        forecast_window.title("Прогнозирование усушки")
        forecast_window.geometry("600x500")
        forecast_window.minsize(500, 400)
        
        # Выбор номенклатуры
        ttk.Label(forecast_window, text="Выберите номенклатуру:").pack(pady=(10, 5))
        nomenclature_var = tk.StringVar()
        nomenclature_combo = ttk.Combobox(forecast_window, textvariable=nomenclature_var, 
                                         values=self.results_data["Номенклатура"].tolist(), 
                                         state="readonly", width=50)
        nomenclature_combo.pack(pady=5)
        
        # Ввод начальной массы
        ttk.Label(forecast_window, text="Начальная масса (кг):").pack(pady=(10, 5))
        initial_mass_var = tk.StringVar(value="100.0")
        initial_mass_entry = ttk.Entry(forecast_window, textvariable=initial_mass_var, width=20)
        initial_mass_entry.pack(pady=5)
        
        # Ввод количества дней
        ttk.Label(forecast_window, text="Количество дней прогноза:").pack(pady=(10, 5))
        days_var = tk.StringVar(value="7")
        days_entry = ttk.Entry(forecast_window, textvariable=days_var, width=20)
        days_entry.pack(pady=5)
        
        # Кнопка для запуска прогноза
        def run_forecast():
            try:
                nomenclature = nomenclature_var.get()
                initial_mass = float(initial_mass_var.get())
                days = int(days_var.get())
                
                if not nomenclature:
                    messagebox.showerror("Ошибка", "Выберите номенклатуру")
                    return
                    
                # Получаем коэффициенты для выбранной номенклатуры
                row = self.results_data[self.results_data["Номенклатура"] == nomenclature].iloc[0]
                coefficients = {
                    'a': row['a'],
                    'b': row['b (день⁻¹)'],
                    'c': row['c']
                }
                
                # Выполняем прогноз
                forecast_result = forecast_shrinkage(coefficients, initial_mass, days)
                
                # Отображаем результаты
                result_window = tk.Toplevel(forecast_window)
                result_window.title(f"Прогноз усушки для {nomenclature}")
                result_window.geometry("700x500")
                result_window.minsize(600, 400)
                
                # Создаем текстовое поле для отображения результатов
                text_frame = ttk.Frame(result_window, padding="10")
                text_frame.pack(expand=True, fill='both')
                
                text_area = scrolledtext.ScrolledText(text_frame, wrap=tk.WORD)
                text_area.pack(expand=True, fill='both')
                
                # Формируем текст с результатами
                result_text = f"ПРОГНОЗ УСУШКИ ДЛЯ НОМЕНКЛАТУРЫ: {nomenclature}\n"
                result_text += f"Начальная масса: {initial_mass} кг\n"
                result_text += f"Период прогноза: {days} дней\n"
                result_text += "=" * 50 + "\n\n"
                
                result_text += "Коэффициенты:\n"
                result_text += f"  a = {coefficients['a']:.3f}\n"
                result_text += f"  b = {coefficients['b']:.3f} (день⁻¹)\n"
                result_text += f"  c = {coefficients['c']:.3f}\n\n"
                
                result_text += "Прогноз по дням:\n"
                for day_data in forecast_result['daily_shrinkage']:
                    result_text += f"  День {day_data['day']:2d}: "
                    result_text += f"усушка {day_data['shrinkage']:8.3f} кг, "
                    result_text += f"накопленная усушка {day_data['cumulative_shrinkage']:8.3f} кг, "
                    result_text += f"остаток {day_data['remaining_mass']:8.3f} кг\n"
                
                result_text += "\n" + "-" * 50 + "\n"
                result_text += f"Общая усушка за {days} дней: {forecast_result['total_shrinkage']:.3f} кг\n"
                result_text += f"Масса после усушки: {forecast_result['final_mass']:.3f} кг\n"
                
                text_area.insert(tk.END, result_text)
                text_area.config(state=tk.DISABLED)
                
            except ValueError as e:
                messagebox.showerror("Ошибка", f"Неверный формат данных: {str(e)}")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Произошла ошибка: {str(e)}")
        
        ttk.Button(forecast_window, text="Рассчитать прогноз", command=run_forecast, 
                  style="Accent.TButton").pack(pady=20)
        
    def compare_coefficients(self):
        """Сравнение коэффициентов"""
        # Открываем окно для выбора файлов
        file_paths = filedialog.askopenfilenames(
            title="Выберите файлы с коэффициентами",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")]
        )
        
        if not file_paths:
            return
            
        try:
            # Выполняем сравнение
            comparison_result = compare_coefficients(list(file_paths))
            
            if comparison_result.empty:
                messagebox.showwarning("Предупреждение", "Не удалось получить данные для сравнения")
                return
            
            # Отображаем результаты
            result_window = tk.Toplevel(self.root)
            result_window.title("Сравнение коэффициентов")
            result_window.geometry("900x600")
            result_window.minsize(800, 500)
            
            # Создаем фрейм для таблицы
            table_frame = ttk.Frame(result_window, padding="10")
            table_frame.pack(expand=True, fill='both')
            
            # Создаем Treeview для отображения результатов
            columns = ("nomenclature", "a_change", "b_change", "c_change", "accuracy_change", 
                      "first_period", "last_period")
            tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=20)
            
            # Определяем заголовки
            tree.heading("nomenclature", text="Номенклатура")
            tree.heading("a_change", text="Изменение A")
            tree.heading("b_change", text="Изменение B")
            tree.heading("c_change", text="Изменение C")
            tree.heading("accuracy_change", text="Изменение точности")
            tree.heading("first_period", text="Первый период")
            tree.heading("last_period", text="Последний период")
            
            # Устанавливаем ширину колонок
            tree.column("nomenclature", width=200)
            tree.column("a_change", width=100)
            tree.column("b_change", width=100)
            tree.column("c_change", width=100)
            tree.column("accuracy_change", width=120)
            tree.column("first_period", width=100)
            tree.column("last_period", width=100)
            
            # Создаем скроллбары
            v_scrollbar = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=tree.yview)
            h_scrollbar = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=tree.xview)
            tree.configure(yscrollcommand=v_scrollbar.set, xscrollcommand=h_scrollbar.set)
            
            # Размещаем элементы
            tree.grid(row=0, column=0, sticky=(tk.N, tk.S, tk.E, tk.W))
            v_scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
            h_scrollbar.grid(row=1, column=0, sticky=(tk.W, tk.E))
            
            # Настраиваем весовые коэффициенты
            table_frame.columnconfigure(0, weight=1)
            table_frame.rowconfigure(0, weight=1)
            
            # Заполняем таблицу данными
            for _, row in comparison_result.iterrows():
                tree.insert("", "end", values=(
                    row["nomenclature"],
                    f"{row['a_change']:.3f}",
                    f"{row['b_change']:.3f}",
                    f"{row['c_change']:.3f}",
                    f"{row['accuracy_change']:.1f}%",
                    row["first_period"],
                    row["last_period"]
                ))
                
        except Exception as e:
            messagebox.showerror("Ошибка", f"Произошла ошибка при сравнении: {str(e)}")
            
    def cluster_nomenclatures(self):
        """Кластеризация номенклатур"""
        if self.results_data is None:
            messagebox.showwarning("Предупреждение", "Результаты еще не рассчитаны")
            return
            
        try:
            # Сохраняем временный файл с результатами для кластеризации
            temp_file = os.path.join(self.project_root, "результаты", "temp_coefficients.csv")
            self.results_data.to_csv(temp_file, index=False)
            
            # Выполняем кластеризацию
            clustering_result = cluster_nomenclatures(temp_file, n_clusters=3)
            
            # Удаляем временный файл
            os.remove(temp_file)
            
            # Отображаем результаты
            result_window = tk.Toplevel(self.root)
            result_window.title("Кластеризация номенклатур")
            result_window.geometry("800x600")
            result_window.minsize(700, 500)
            
            # Создаем текстовое поле для отображения результатов
            text_frame = ttk.Frame(result_window, padding="10")
            text_frame.pack(expand=True, fill='both')
            
            text_area = scrolledtext.ScrolledText(text_frame, wrap=tk.WORD)
            text_area.pack(expand=True, fill='both')
            
            # Формируем текст с результатами
            result_text = "РЕЗУЛЬТАТЫ КЛАСТЕРИЗАЦИИ НОМЕНКЛАТУР\n"
            result_text += "=" * 50 + "\n\n"
            
            result_text += f"Общее количество номенклатур: {clustering_result['total_nomenclatures']}\n"
            result_text += f"Количество кластеров: {clustering_result['n_clusters']}\n\n"
            
            for cluster_id, cluster_info in clustering_result['clusters'].items():
                result_text += f"КЛАСТЕР {cluster_id}:\n"
                result_text += f"  Количество номенклатур: {cluster_info['count']}\n"
                result_text += f"  Центр кластера:\n"
                result_text += f"    a = {cluster_info['center']['a']:.3f}\n"
                result_text += f"    b = {cluster_info['center']['b']:.3f}\n"
                result_text += f"    c = {cluster_info['center']['c']:.3f}\n"
                result_text += "  Номенклатуры:\n"
                for nom in cluster_info['nomenclatures'][:10]:  # Показываем первые 10
                    result_text += f"    - {nom}\n"
                if len(cluster_info['nomenclatures']) > 10:
                    result_text += f"    ... и еще {len(cluster_info['nomenclatures']) - 10} позиций\n"
                result_text += "\n" + "-" * 30 + "\n\n"
                
            text_area.insert(tk.END, result_text)
            text_area.config(state=tk.DISABLED)
            
        except Exception as e:
            messagebox.showerror("Ошибка", f"Произошла ошибка при кластеризации: {str(e)}")

def main():
    root = tk.Tk()
    app = ShrinkageCalculatorGUI(root)
    root.mainloop()

if __name__ == "__main__":
    main()